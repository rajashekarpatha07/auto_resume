"""Professional ATS-friendly DOCX resume builder using python-docx."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from Agent.schemas.resume import (
    EducationSchema,
    ProjectsSchema,
    ResumeHeaderSchema,
    SkillsSchema,
    SummarySchema,
)

# ──────────────── Style Constants ────────────────
FONT_NAME = "Calibri"
NAME_SIZE = Pt(20)
HEADLINE_SIZE = Pt(11)
CONTACT_SIZE = Pt(9)
SECTION_HEADER_SIZE = Pt(12)
BODY_SIZE = Pt(10.5)
SUBHEADER_SIZE = Pt(11)

DARK_BLUE = RGBColor(0, 51, 102)
DARK_GRAY = RGBColor(64, 64, 64)
MEDIUM_GRAY = RGBColor(100, 100, 100)
BLACK = RGBColor(0, 0, 0)

MARGIN_TOP = Inches(0.5)
MARGIN_BOTTOM = Inches(0.5)
MARGIN_LEFT = Inches(0.6)
MARGIN_RIGHT = Inches(0.6)

DEFAULT_SECTION_ORDER = ["summary", "skills", "projects", "education"]


# ──────────────── Helpers ────────────────


def _set_spacing(para, before=0, after=0):
    """Set paragraph spacing in points."""
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def _add_bottom_border(paragraph, color="003366", size="6"):
    """Add a thin bottom border line under a paragraph."""
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _styled_run(para, text, *, bold=False, italic=False, size=BODY_SIZE, color=BLACK):
    """Add a styled run to a paragraph."""
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = size
    run.font.name = FONT_NAME
    run.font.color.rgb = color
    return run


def _section_header(doc, text):
    """Add a dark-blue, uppercase section header with a bottom border."""
    para = doc.add_paragraph()
    _set_spacing(para, before=10, after=4)
    _styled_run(para, text.upper(), bold=True, size=SECTION_HEADER_SIZE, color=DARK_BLUE)
    _add_bottom_border(para)
    return para


# ──────────────── Section Builders ────────────────


def _add_header(doc, header: ResumeHeaderSchema):
    """Name, headline, contact line — centred."""
    # Name
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(name_para, before=0, after=0)
    _styled_run(name_para, header.full_name, bold=True, size=NAME_SIZE)

    # Headline
    hl_para = doc.add_paragraph()
    hl_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(hl_para, before=0, after=2)
    _styled_run(hl_para, header.headline, size=HEADLINE_SIZE, color=MEDIUM_GRAY)

    # Contact
    ct_para = doc.add_paragraph()
    ct_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(ct_para, before=0, after=6)
    _styled_run(
        ct_para,
        " │ ".join(header.contact_items),
        size=CONTACT_SIZE,
        color=DARK_GRAY,
    )


def _add_summary(doc, summary: SummarySchema):
    _section_header(doc, "Summary")
    para = doc.add_paragraph()
    _set_spacing(para, before=2, after=4)
    _styled_run(para, summary.summary)


def _add_skills(doc, skills: SkillsSchema):
    _section_header(doc, "Skills")
    for cat in skills.categories:
        para = doc.add_paragraph()
        _set_spacing(para, before=1, after=1)
        _styled_run(para, f"{cat.category}: ", bold=True)
        _styled_run(para, ", ".join(cat.skills))


def _add_projects(doc, projects: ProjectsSchema):
    _section_header(doc, "Projects")
    for p in projects.projects:
        # Project title
        title_para = doc.add_paragraph()
        _set_spacing(title_para, before=4, after=0)
        _styled_run(title_para, p.name, bold=True, size=SUBHEADER_SIZE)
        if p.link:
            _styled_run(title_para, f"  |  {p.link}", size=Pt(9), color=MEDIUM_GRAY)

        # Tech stack (italic)
        tech_para = doc.add_paragraph()
        _set_spacing(tech_para, before=0, after=1)
        _styled_run(
            tech_para,
            ", ".join(p.tech_stack),
            italic=True,
            size=Pt(10),
            color=DARK_GRAY,
        )

        # Bullets
        for bullet in p.bullets:
            bp = doc.add_paragraph(style="List Bullet")
            _set_spacing(bp, before=0, after=0)
            _styled_run(bp, bullet)


def _add_education(doc, education: EducationSchema):
    _section_header(doc, "Education")
    for edu in education.education:
        inst_para = doc.add_paragraph()
        _set_spacing(inst_para, before=4, after=0)
        _styled_run(inst_para, edu.institution, bold=True, size=SUBHEADER_SIZE)

        deg_para = doc.add_paragraph()
        _set_spacing(deg_para, before=0, after=1)
        _styled_run(deg_para, edu.degree)
        meta = [x for x in [edu.duration, edu.grade, edu.location] if x]
        if meta:
            _styled_run(deg_para, f"  |  {' | '.join(meta)}", size=Pt(10), color=MEDIUM_GRAY)

        for h in edu.highlights:
            hp = doc.add_paragraph(style="List Bullet")
            _set_spacing(hp, before=0, after=0)
            _styled_run(hp, h)

    if education.extracurriculars:
        extra_para = doc.add_paragraph()
        _set_spacing(extra_para, before=6, after=2)
        _styled_run(extra_para, "Extracurricular Activities", bold=True)
        for a in education.extracurriculars:
            ap = doc.add_paragraph(style="List Bullet")
            _set_spacing(ap, before=0, after=0)
            _styled_run(ap, a)


# ──────────────── Public API ────────────────


def build_resume_docx(
    doc_path: Path,
    header: ResumeHeaderSchema,
    summary: SummarySchema,
    skills: SkillsSchema,
    projects: ProjectsSchema,
    education: EducationSchema,
    section_order: list[str] | None = None,
) -> None:
    """Build a professional, ATS-friendly single-column DOCX resume.

    ``section_order`` controls the order of the body sections (from the
    optimization node).  Falls back to the default order if not supplied.
    """
    if section_order is None:
        section_order = DEFAULT_SECTION_ORDER

    doc = Document()

    # Global defaults
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = BODY_SIZE

    for sec in doc.sections:
        sec.top_margin = MARGIN_TOP
        sec.bottom_margin = MARGIN_BOTTOM
        sec.left_margin = MARGIN_LEFT
        sec.right_margin = MARGIN_RIGHT

    # Header is always first
    _add_header(doc, header)

    # Body sections in optimised order
    builders = {
        "summary": lambda: _add_summary(doc, summary),
        "skills": lambda: _add_skills(doc, skills),
        "projects": lambda: _add_projects(doc, projects),
        "education": lambda: _add_education(doc, education),
    }
    for name in section_order:
        if name in builders:
            builders[name]()

    doc.save(str(doc_path))
